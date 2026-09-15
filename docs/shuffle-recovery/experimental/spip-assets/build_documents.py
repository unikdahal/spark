#!/usr/bin/env python3
# Licensed to the Apache Software Foundation (ASF) under one or more
# contributor license agreements. See the NOTICE file distributed with
# this work for additional information regarding copyright ownership.
# The ASF licenses this file to You under the Apache License, Version 2.0
# (the "License"); you may not use this file except in compliance with
# the License. You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Render the Markdown discussion package to PDF and editable Word documents.

Install spip-assets/requirements.txt in a temporary virtual environment, then run
this file. PDFs and DOCX use the same parsed Markdown and explicit page breaks.
"""

from html import escape
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET

import pymupdf as fitz
import markdown
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    Image, PageBreak, Paragraph, Preformatted, SimpleDocTemplate,
    Spacer, Table, TableStyle,
)

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "spip-assets"
PUBLIC = (
    "https://github.com/unikdahal/spark/blob/completed-shuffle-reuse-discussion/"
    "docs/shuffle-recovery/experimental/"
)
INK = "233248"
BLUE = "235B88"
LIGHT = "EDF2F7"
WIDTH = A4[0] - 104


def external(href):
    if href.startswith(("https://", "http://")):
        return href
    return PUBLIC + href


def diagram(name, labels):
    """Draw four readable process boxes; retain both vector PDF and Word PNG."""
    stream = BytesIO()
    c = canvas.Canvas(stream, pagesize=(500, 112))
    box_width = 113
    for i, (title, lines) in enumerate(labels):
        x = 3 + i * 128
        c.setFillColor(colors.HexColor("#" + LIGHT))
        c.setStrokeColor(colors.HexColor("#BAC7D5"))
        c.roundRect(x, 16, box_width, 88, 5, stroke=1, fill=1)
        c.setFillColor(colors.HexColor("#" + INK))
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(x + box_width / 2, 84, title)
        c.setFont("Helvetica", 9)
        for j, text in enumerate(lines):
            c.drawCentredString(x + box_width / 2, 66 - j * 13, text)
        if i < 3:
            c.setStrokeColor(colors.HexColor("#" + BLUE))
            c.setLineWidth(1.2)
            start = x + box_width + 2
            c.line(start, 59, start + 11, 59)
            c.line(start + 7, 63, start + 11, 59)
            c.line(start + 7, 55, start + 11, 59)
    c.showPage()
    c.save()
    pdf = stream.getvalue()
    (ASSETS / f"{name}.pdf").write_bytes(pdf)
    with fitz.open(stream=pdf, filetype="pdf") as doc:
        doc[0].get_pixmap(matrix=fitz.Matrix(3, 3)).save(ASSETS / f"{name}.png")


def inline(element):
    text = escape(element.text or "")
    for child in element:
        content = inline(child)
        if child.tag == "strong":
            content = f"<b>{content}</b>"
        elif child.tag == "em":
            content = f"<i>{content}</i>"
        elif child.tag == "code":
            content = f'<font name="Courier" size="8.5">{content}</font>'
        elif child.tag == "a":
            url = escape(external(child.attrib["href"]), quote=True)
            content = f'<a href="{url}" color="#{BLUE}">{content}</a>'
        elif child.tag == "br":
            content = "<br/>"
        text += content + escape(child.tail or "")
    return text


STYLES = {
    "p": ParagraphStyle(
        "Body", fontName="Helvetica", fontSize=10.5, leading=14.1,
        textColor=colors.HexColor("#20252C"), spaceAfter=8,
    ),
    "h1": ParagraphStyle(
        "Title", fontName="Helvetica-Bold", fontSize=23, leading=27,
        textColor=colors.HexColor("#" + INK), spaceAfter=16,
    ),
    "h2": ParagraphStyle(
        "Section", fontName="Helvetica-Bold", fontSize=15, leading=19,
        textColor=colors.HexColor("#" + INK), spaceBefore=7, spaceAfter=11,
        keepWithNext=True,
    ),
    "h3": ParagraphStyle(
        "Subsection", fontName="Helvetica-Bold", fontSize=11.5, leading=15,
        textColor=colors.HexColor("#" + INK), spaceBefore=6, spaceAfter=7,
        keepWithNext=True,
    ),
    "cell": ParagraphStyle(
        "Cell", fontName="Helvetica", fontSize=9.4, leading=12.2,
        spaceAfter=0,
    ),
    "code": ParagraphStyle(
        "Code", fontName="Courier", fontSize=8.3, leading=11,
        backColor=colors.HexColor("#F3F5F7"), borderPadding=8,
        spaceBefore=5, spaceAfter=12,
    ),
}


def table_weights(element):
    header = element.find(".//tr")
    n = len(header)
    if "".join(header[0].itertext()) == "Surface":
        return [0.49, 0.51]
    weights = {2: [0.30, 0.70], 3: [0.35, 0.17, 0.48],
               4: [0.42, 0.17, 0.16, 0.25], 5: [0.26, 0.12, 0.17, 0.18, 0.27]}
    return weights.get(n, [1 / n] * n)


def pdf_blocks(element):
    if element.tag in ("h1", "h2", "h3"):
        return [Paragraph(inline(element), STYLES[element.tag])]
    if element.tag == "p":
        img = element.find("img")
        if img is not None:
            path = ROOT / img.attrib["src"]
            return [Image(str(path), width=WIDTH, height=WIDTH * 112 / 500),
                    Spacer(1, 7)]
        return [Paragraph(inline(element), STYLES["p"])]
    if element.tag == "pre":
        return [Preformatted("".join(element.itertext()).rstrip(), STYLES["code"])]
    if element.tag in ("ul", "ol"):
        result = []
        for i, item in enumerate(element):
            prefix = f"{i + 1}. " if element.tag == "ol" else "&#8226; "
            result.append(Paragraph(prefix + inline(item), STYLES["p"]))
        return result
    if element.tag == "table":
        rows = []
        for row in element.findall(".//tr"):
            rows.append([
                Paragraph(inline(cell), STYLES["cell"]) for cell in row
            ])
        widths = [WIDTH * w for w in table_weights(element)]
        table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + LIGHT)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFBFC")]),
            ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor("#BBC8D6")),
            ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor("#DDE3E9")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return [table, Spacer(1, 10)]
    return []


def word_inline(paragraph, element, bold=False, italic=False, code=False):
    def add(text):
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
    add(element.text)
    for child in element:
        if child.tag == "a":
            link = OxmlElement("w:hyperlink")
            rel = paragraph.part.relate_to(external(child.attrib["href"]), RT.HYPERLINK,
                                          is_external=True)
            link.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), BLUE)
            props.append(color)
            run.append(props)
            text = OxmlElement("w:t")
            text.text = "".join(child.itertext())
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        elif child.tag == "br":
            paragraph.add_run().add_break()
        else:
            word_inline(paragraph, child, bold or child.tag == "strong",
                        italic or child.tag == "em", code or child.tag == "code")
        add(child.tail)


def word_blocks(doc, element):
    if element.tag in ("h1", "h2", "h3", "p"):
        img = element.find("img")
        if img is not None:
            doc.add_picture(str(ROOT / img.attrib["src"]), width=Inches(WIDTH / 72))
            return
        style = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2"}.get(element.tag)
        word_inline(doc.add_paragraph(style=style), element)
    elif element.tag == "pre":
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run("".join(element.itertext()).rstrip())
        run.font.name, run.font.size = "Consolas", Pt(8.3)
    elif element.tag in ("ol", "ul"):
        for i, item in enumerate(element):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.12)
            p.add_run(f"{i + 1}. " if element.tag == "ol" else "\u2022 ")
            word_inline(p, item)
    elif element.tag == "table":
        rows = element.findall(".//tr")
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        table.autofit = False
        widths = [Inches(WIDTH * w / 72) for w in table_weights(element)]
        for column, width in zip(table.columns, widths):
            column.width = width
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, width in zip(cells, widths):
                cell.width = width
            tr_props = table.rows[-1]._tr.get_or_add_trPr()
            tr_props.append(OxmlElement("w:cantSplit"))
            if i == 0:
                tr_props.append(OxmlElement("w:tblHeader"))
            for cell, content in zip(cells, row):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                word_inline(p, content, bold=i == 0)
                for run in p.runs:
                    run.font.size = Pt(9.3)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), LIGHT if i == 0 else "FFFFFF")
                cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)


def render(stem, short_title):
    source = (ROOT / f"{stem}.md").read_text()
    pages = source.split("<!-- pagebreak -->")
    trees = [ET.fromstring("<root>" + markdown.markdown(
        page, extensions=["tables", "fenced_code"]
    ) + "</root>") for page in pages]
    story = []
    for i, tree in enumerate(trees):
        if i:
            story.append(PageBreak())
        for element in tree:
            story.extend(pdf_blocks(element))

    def page_decor(c, doc):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#657286"))
        c.drawString(52, A4[1] - 31, short_title)
        c.drawRightString(A4[0] - 52, A4[1] - 31, "DISCUSSION DRAFT")
        c.setStrokeColor(colors.HexColor("#D5DEE7"))
        c.line(52, 43, A4[0] - 52, 43)
        c.drawString(52, 29, "Unik Dahal | Completed-shuffle reuse")
        c.drawRightString(A4[0] - 52, 29, str(doc.page))
        c.restoreState()

    pdf = SimpleDocTemplate(
        str(ROOT / f"{stem}.pdf"), pagesize=A4, leftMargin=52, rightMargin=52,
        topMargin=53, bottomMargin=57, title=short_title, author="Unik Dahal",
    )
    pdf.build(story, onFirstPage=page_decor, onLaterPages=page_decor)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(A4[0] / 72), Inches(A4[1] / 72)
    section.left_margin = section.right_margin = Inches(52 / 72)
    section.top_margin, section.bottom_margin = Inches(53 / 72), Inches(57 / 72)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(8)
    for name, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 11.5)]:
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_after = Pt(10)
    header = section.header.paragraphs[0]
    header.add_run(short_title + "  |  Discussion draft").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run("Unik Dahal | Completed-shuffle reuse  |  Page ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = short_title
    doc.core_properties.author = "Unik Dahal"
    doc.core_properties.subject = "Completed-shuffle reuse proposal"
    for i, tree in enumerate(trees):
        if i:
            doc.add_page_break()
        for element in tree:
            word_blocks(doc, element)
    doc.save(ROOT / f"{stem}.docx")
    with fitz.open(ROOT / f"{stem}.pdf") as rendered:
        print(f"{stem}: {len(rendered)} PDF pages; {len(pages)} intended sections")
        for i, page in enumerate(rendered):
            headings = page.get_text().splitlines()[:3]
            print(i + 1, " | ".join(headings))
        if len(rendered) != len(pages):
            raise RuntimeError("Layout overflow: inspect pages before publishing")


if __name__ == "__main__":
    diagram("architecture", [
        ("Current source", ["Actual planned read", "Certificate + splits"]),
        ("Spark SQL", ["Canonical identity", "Supported producer"]),
        ("Manifest + provider", ["Full identity match", "Native read claim"]),
        ("Spark scheduler", ["Local adoption", "Native handle read"]),
    ])
    diagram("publication", [
        ("Accepted maps", ["Ordered task winners", "Current tracker check"]),
        ("Provider seal", ["Native descriptor", "Recheck selection"]),
        ("Immutable store", ["Body then index", "Full identity record"]),
        ("Replacement", ["Certify current read", "Find and claim"]),
    ])
    diagram("failure", [
        ("Native read fails", ["Binding + shuffle ID", "FetchFailed event"]),
        ("Fence binding", ["Clear adopted status", "Advance epoch"]),
        ("Scheduler retry", ["Whole-stage marker", "Existing recovery"]),
        ("Fresh execution", ["Ordinary producer", "Recomputed result"]),
    ])
    render("spip-proposal", "Completed-shuffle reuse | SPIP")
    render("spip-design-evidence", "Completed-shuffle reuse | Design notes")
